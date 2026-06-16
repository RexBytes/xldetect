"""Turning ``openpyxl`` worksheets into :class:`~xldetect.grid.Grid` objects.

This is the only module that touches ``openpyxl`` directly. It loads a workbook
with ``data_only=True`` (cached formula *values*, never formula text -- formula
evaluation is out of scope) and ``read_only=False`` (so cell styles are
available for header detection). Merged ranges are recorded and their areas are
filled with the anchor (top-left) value, so a merge counts as occupied across
its whole span. Cells that are ``None`` or whitespace-only strings are treated
as blank and never enter the grid.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from .grid import CellStyle, Grid


def _extract_style(cell) -> CellStyle:
    """Read the four header-relevant formatting cues from an ``openpyxl`` cell."""
    font = cell.font
    bold = bool(font and font.bold)
    italic = bool(font and font.italic)

    fill = cell.fill
    filled = bool(fill is not None and getattr(fill, "patternType", None) not in (None, "none"))

    border = cell.border
    bordered = False
    if border is not None:
        bordered = any(
            getattr(border, side) is not None and getattr(border, side).style
            for side in ("left", "right", "top", "bottom")
        )
    return CellStyle(bold=bold, filled=filled, bordered=bordered, italic=italic)


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def build_grid(ws) -> Grid:
    """Build a :class:`Grid` from a live ``openpyxl`` worksheet."""
    values: dict[tuple[int, int], object] = {}
    styles: dict[tuple[int, int], CellStyle] = {}

    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if _is_blank(v):
                continue
            values[(cell.row, cell.column)] = v
            styles[(cell.row, cell.column)] = _extract_style(cell)

    merged: list[tuple[int, int, int, int]] = []
    for rng in ws.merged_cells.ranges:
        merged.append((rng.min_row, rng.min_col, rng.max_row, rng.max_col))

    # Fill merged areas with the anchor value so a merge is occupied across its span.
    for mr0, mc0, mr1, mc1 in merged:
        anchor = values.get((mr0, mc0))
        if anchor is None:
            continue
        anchor_style = styles.get((mr0, mc0), CellStyle())
        for r in range(mr0, mr1 + 1):
            for c in range(mc0, mc1 + 1):
                values.setdefault((r, c), anchor)
                styles.setdefault((r, c), anchor_style)

    return Grid(values=values, styles=styles, merged=merged, sheet=ws.title)


def load_grids(
    path: str | Path, sheets: list[str] | None = None
) -> dict[str, Grid]:
    """Load some or all worksheets of a workbook as :class:`Grid` objects.

    Args:
        path: Path to an ``.xlsx`` workbook.
        sheets: Optional list of sheet names to load; ``None`` loads all sheets
            in workbook order. An unknown name raises :class:`ValueError`.

    Returns:
        A dict mapping sheet name to :class:`Grid`, preserving the requested
        (or workbook) order.

    Raises:
        ValueError: if the file is missing, not a valid ``.xlsx`` workbook
            (corrupt, wrong format, not a zip), or names an unknown sheet. This
            normalises openpyxl's lower-level exceptions so every caller -- the
            CLI included -- sees a single, user-facing error type.
    """
    try:
        wb = openpyxl.load_workbook(filename=str(path), data_only=True, read_only=False)
    except FileNotFoundError as exc:
        raise ValueError(f"File not found: {path}") from exc
    except (InvalidFileException, zipfile.BadZipFile) as exc:
        raise ValueError(
            f"Not a readable .xlsx workbook ({path}): {exc}"
        ) from exc
    except OSError as exc:
        raise ValueError(f"Could not open workbook ({path}): {exc}") from exc
    try:
        names = list(sheets) if sheets is not None else list(wb.sheetnames)
        out: dict[str, Grid] = {}
        for name in names:
            if name not in wb.sheetnames:
                raise ValueError(
                    f"Sheet {name!r} not found. Available: {wb.sheetnames}"
                )
            out[name] = build_grid(wb[name])
        return out
    finally:
        wb.close()
